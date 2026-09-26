from io import BytesIO
import pandas as pd
import openpyxl

from google_sheets_roster import parse_dob, last4_from_nric
from reference_lists import ENTRY_HEADERS
from .formatting import normalize_header, gender_to_code, code_to_gender_display
from .validation import normalize_ic_last4, normalize_email

def export_entries_to_excel(header_info: dict, entries: pd.DataFrame) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entry Form"

    pairs = [
        ("Team Name", "team_name"),
        ("Billing Contact Name", "billing_name"),
        ("Billing Email", "billing_email"),
        ("Charge Code", "charge_code"),
        ("P/O to be sent", "po_to_be_sent"),
    ]
    for row_num, (label, key) in enumerate(pairs, start=1):
        ws.cell(row_num, 1).value = label
        ws.cell(row_num, 2).value = header_info.get(key, "")

    header_row = 7
    for col, heading in enumerate(ENTRY_HEADERS, start=1):
        ws.cell(header_row, col).value = heading

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()

def sheet_df_to_entries(df: pd.DataFrame) -> list[dict]:
    if df is None or df.empty:
        return []

    col_map = {normalize_header(c): c for c in df.columns}

    def get(row, key, default=""):
        column = col_map.get(key)
        return row.get(column, default) if column else default

    entries = []
    for _, row in df.iterrows():
        # Admin withdrawals are soft deletes. Keep the historical OUTPUT row
        # for traceability but exclude it from operational entry exports.
        is_deleted_raw = str(get(row, "is_deleted", "") or "").strip().casefold()
        entry_status = str(get(row, "entry_status", "") or "").strip().upper()
        if is_deleted_raw in {"true", "1", "yes", "y"} or entry_status == "WITHDRAWN":
            continue

        first_name = str(get(row, "first_name", "") or get(row, "firstname", "") or get(row, "first", "")).strip()
        other_name = str(get(row, "other_name", "") or get(row, "othername", "")).strip()
        last_name = str(get(row, "last_name", "") or get(row, "lastname", "") or get(row, "last", "")).strip()

        gender_raw = str(get(row, "gender", "")).strip()
        gcode = gender_to_code(gender_raw)
        gender = code_to_gender_display(gcode) or gender_raw

        dob_raw = get(row, "birth_date", "") or get(row, "dob", "") or get(row, "date_of_birth", "")
        try:
            birth_date = parse_dob(dob_raw)
        except Exception:
            birth_date = None

        nric_raw = str(get(row, "nric", "") or "").strip()
        ic_last4 = str(get(row, "ic_last4", "") or "").strip()
        if not ic_last4 and nric_raw:
            ic_last4 = last4_from_nric(nric_raw)
        ic_last4 = normalize_ic_last4(ic_last4)

        nationality = str(get(row, "nationality", "")).strip()
        singapore_pr = str(get(row, "singapore_pr", "") or get(row, "sg_pr", "") or "").strip()
        if singapore_pr.casefold() in ("true", "1", "y", "yes"):
            singapore_pr = "Yes"
        elif singapore_pr.casefold() in ("false", "0", "n", "no"):
            singapore_pr = "No"
        else:
            singapore_pr = singapore_pr or "No"

        unique_id = str(get(row, "unique_id", "")).strip()
        team_name = str(get(row, "team_name", "")).strip()
        team_code = str(get(row, "team_code", "")).strip()
        event = str(get(row, "event", "")).strip()
        event_code = str(get(row, "event_code", "")).strip()
        season_best = str(get(row, "season_best", "") or "").strip()
        parq = str(get(row, "parq", "") or "").strip()
        charge_code = str(get(row, "charge_code", "")).strip()
        po_to_be_sent = str(get(row, "po_to_be_sent", "")).strip()
        email = normalize_email(str(get(row, "email", "")).strip())
        contact_number = str(get(row, "contact_number", "") or get(row, "contact", "")).strip()

        full_name_sheet = str(get(row, "full_name", "") or get(row, "full", "") or "").strip()
        name = " ".join(p for p in [first_name, other_name, last_name] if p).strip()
        if full_name_sheet:
            name = name or full_name_sheet

        if not any([name, unique_id, team_code, team_name, email, contact_number, ic_last4]):
            continue

        entries.append({
            "name": name,
            "full_name": full_name_sheet or name,
            "name_passport": str(get(row, "name_passport", "") or get(row, "name_as_per_nric_passport", "") or "").strip() or (full_name_sheet or name),
            "first_name": first_name,
            "other_name": other_name,
            "last_name": last_name,
            "gender": gender if gender in ("Male", "Female", "M", "F") else "",
            "birth_date": birth_date,
            "ic_last4": ic_last4,
            "nationality": nationality,
            "singapore_pr": singapore_pr,
            "unique_id": unique_id,
            "team_name": team_name,
            "team_code": team_code,
            "event": event,
            "event_code": event_code,
            "charge_code": charge_code,
            "po_to_be_sent": po_to_be_sent if po_to_be_sent in ("Yes", "No") else "",
            "season_best": season_best,
            "parq": parq if parq in ("Y", "N") else "Y",
            "emergency_contact_name": str(get(row, "emergency_contact_name", "") or "").strip(),
            "emergency_contact_number": str(get(row, "emergency_contact_number", "") or "").strip(),
            "coach_full_name": str(get(row, "coach_full_name", "") or "").strip(),
            "email": email,
            "contact_number": contact_number,
        })

    return entries

def build_semicolon_export_from_output_sheet(sheet_df: pd.DataFrame, record_type: str = "I") -> str:
    """Build the established Hy-Tek I/E semicolon export from legacy OUTPUT.

    Phase 6A keeps this compatibility entry point because the private/admin site
    already imports it. New exports should use the transaction-backed builder
    in :mod:`signup.hytek_export`.
    """
    from .hytek_export import build_legacy_output_hytek_export

    return build_legacy_output_hytek_export(sheet_df, record_type=record_type)
